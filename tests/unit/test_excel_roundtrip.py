"""
Excel round-trip fidelity tests (v3.0.13).

A Mediatools -> Planmatic round-trip test found that JSON -> Excel -> JSON was
lossless only at line-item level. The causes were label/format mismatches
between the exporter and importer, and an importer that rebuilt
``metric_formulas`` from scratch. These tests pin each fix, plus
compatibility with workbooks exported by v3.0.12 and earlier.

Tests that need formula results run the workbook through LibreOffice
(``soffice --convert-to xlsx``) and are skipped when it is not installed.
openpyxl cannot evaluate formulas, so these tests need a real spreadsheet engine.
"""

import copy
import json
import shutil
import subprocess
import warnings

import openpyxl
import pytest

from mediaplanpy.excel.exporter import export_to_excel
from mediaplanpy.excel.importer import (
    ExcelFormulaCacheWarning,
    _merge_metric_formulas,
    _parse_list_field,
    import_from_excel,
)
from mediaplanpy.models import MediaPlan

SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
needs_soffice = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")

# Most tests import a workbook that was never recalculated, which is exactly
# what the B7 warning is for; it is asserted explicitly in its own tests.
pytestmark = pytest.mark.filterwarnings(
    "ignore::mediaplanpy.excel.importer.ExcelFormulaCacheWarning"
)


def _plan_dict():
    return {
        "meta": {
            "id": "mediaplan_rt",
            "schema_version": "v3.0",
            "name": "Round trip",
            "created_by_name": "tester",
            "created_at": "2026-01-01T00:00:00Z",
            "custom_properties": {"mt_plan_id": "P-1", "calendar": {"weeks": [1, 2]}},
        },
        "campaign": {
            "id": "campaign_rt",
            "name": "Campaign",
            "start_date": "2026-01-01",
            "end_date": "2026-03-31",
            "budget_total": 50000,
            "custom_properties": {"mt_client_id": "C-9"},
            "target_locations": [
                {
                    "name": "Spot markets",
                    "location_type": "DMA",
                    "location_list": ["New York", "Los Angeles, CA"],
                    "exclusion_type": "DMA",
                    "exclusion_list": ["Boston"],
                }
            ],
        },
        "lineitems": [
            {
                "id": "li_full",
                "name": "Full",
                "start_date": "2026-01-01",
                "end_date": "2026-02-01",
                "cost_total": 10000,
                "cost_media": 8000,
                "cost_currency": "USD",
                "metric_impressions": 2000000,
                "metric_clicks": 20000,
                "metric_audience_size": 1500000,
                "metric_formulas": {
                    "metric_impressions": {
                        "formula_type": "cost_per_unit",
                        "base_metric": "cost_total",
                        "coefficient": 0.005,
                        "comments": "GRPsAndPop methodology",
                    },
                    "metric_clicks": {
                        "formula_type": "conversion_rate",
                        "base_metric": "metric_impressions",
                        "coefficient": 0.01,
                        "parameter3": 7.0,
                    },
                    "metric_audience_size": {"formula_type": "constant", "coefficient": 1500000},
                },
                "custom_properties": {"mt_row": 1},
            },
            {
                # Has none of the cost breakdown / metric fields the other row has
                "id": "li_sparse",
                "name": "Sparse",
                "start_date": "2026-01-01",
                "end_date": "2026-02-01",
                "cost_total": 500,
            },
        ],
        "dictionary": {
            "lineitem_custom_dimensions": {
                "dim_custom1": {"status": "enabled", "caption": "Market"}
            },
        },
    }


def _export(tmp_path, plan=None):
    media_plan = MediaPlan.from_dict(copy.deepcopy(plan or _plan_dict()))
    path = tmp_path / "plan.xlsx"
    export_to_excel(media_plan, str(path))
    return path


def _recalculate(path, tmp_path):
    """Open and save the workbook in LibreOffice so formula results get cached."""
    out_dir = tmp_path / "recalc"
    subprocess.run(
        [
            SOFFICE,
            "--headless",
            "--calc",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out_dir),
            str(path),
        ],
        check=True,
        capture_output=True,
        timeout=120,
    )
    return out_dir / path.name


def _roundtrip(tmp_path, recalculate=False):
    path = _export(tmp_path)
    if recalculate:
        path = _recalculate(path, tmp_path)
    return MediaPlan.from_dict(import_from_excel(str(path))).to_dict()


def _lineitem(plan, lineitem_id):
    return next(li for li in plan["lineitems"] if li["id"] == lineitem_id)


class TestPlanAndCampaignFields:

    def test_meta_custom_properties_survive(self, tmp_path):  # B1
        plan = _roundtrip(tmp_path)
        assert plan["meta"]["custom_properties"] == _plan_dict()["meta"]["custom_properties"]

    def test_campaign_custom_properties_survive(self, tmp_path):  # B2
        plan = _roundtrip(tmp_path)
        assert plan["campaign"]["custom_properties"] == {"mt_client_id": "C-9"}

    def test_location_lists_survive_including_commas(self, tmp_path):  # B3
        location = _roundtrip(tmp_path)["campaign"]["target_locations"][0]
        assert location["location_list"] == ["New York", "Los Angeles, CA"]
        assert location["exclusion_list"] == ["Boston"]

    def test_location_lists_exported_as_json(self, tmp_path):
        sheet = openpyxl.load_workbook(_export(tmp_path))["Target Locations"]
        headers = [c.value for c in sheet[1]]
        assert "Location List (JSON)" in headers
        assert "Exclusion List (JSON)" in headers
        cell = sheet.cell(row=2, column=headers.index("Location List (JSON)") + 1).value
        assert json.loads(cell) == ["New York", "Los Angeles, CA"]


class TestLineItemFields:

    def test_cost_currency_survives(self, tmp_path):  # B5
        assert _lineitem(_roundtrip(tmp_path), "li_full")["cost_currency"] == "USD"

    def test_custom_properties_survive(self, tmp_path):
        assert _lineitem(_roundtrip(tmp_path), "li_full")["custom_properties"] == {"mt_row": 1}

    @pytest.mark.parametrize("recalculate", [False, pytest.param(True, marks=needs_soffice)])
    def test_metric_formula_extras_survive(self, tmp_path, recalculate):  # B4, B10
        formulas = _lineitem(_roundtrip(tmp_path, recalculate), "li_full")["metric_formulas"]
        assert formulas["metric_impressions"]["comments"] == "GRPsAndPop methodology"
        assert formulas["metric_clicks"]["parameter3"] == 7.0
        assert formulas["metric_audience_size"]["formula_type"] == "constant"
        assert formulas["metric_audience_size"]["coefficient"] == 1500000

    @needs_soffice
    def test_edited_coefficient_wins_over_json_column(self, tmp_path):
        """The sheet is where users edit; its coefficient must beat the stale JSON column."""
        path = _recalculate(_export(tmp_path), tmp_path)
        wb = openpyxl.load_workbook(path)
        sheet = wb["Line Items"]
        headers = [c.value for c in sheet[1]]
        cost_col = headers.index("Cost Total") + 1
        sheet.cell(row=2, column=cost_col, value=20000)  # double spend, CPM unchanged
        wb.save(path)
        path = _recalculate(path, tmp_path / "second")

        formula = _lineitem(MediaPlan.from_dict(import_from_excel(str(path))).to_dict(), "li_full")[
            "metric_formulas"
        ]["metric_impressions"]
        assert formula["coefficient"] == pytest.approx(0.005)
        assert formula["comments"] == "GRPsAndPop methodology"


@needs_soffice
class TestAbsentStaysAbsent:  # B14

    def test_sparse_lineitem_gains_no_zero_fields(self, tmp_path):
        sparse = _lineitem(_roundtrip(tmp_path, recalculate=True), "li_sparse")
        original = _lineitem(_plan_dict(), "li_sparse")
        added = {k: v for k, v in sparse.items() if k not in original and v not in (None, {}, [])}
        assert added == {}

    def test_present_values_still_computed(self, tmp_path):
        full = _lineitem(_roundtrip(tmp_path, recalculate=True), "li_full")
        assert full["cost_media"] == pytest.approx(8000)
        assert full["metric_impressions"] == pytest.approx(2000000)
        assert full["metric_clicks"] == pytest.approx(20000)


class TestAbsentStaysAbsentExport:  # B14, export side (no recalculation needed)

    def test_absent_cost_pct_blank_and_formula_guarded(self, tmp_path):
        sheet = openpyxl.load_workbook(_export(tmp_path))["Line Items"]
        headers = [c.value for c in sheet[1]]
        sparse_row = next(
            r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value == "li_sparse"
        )
        pct = sheet.cell(sparse_row, headers.index("Cost Media %") + 1).value
        formula = sheet.cell(sparse_row, headers.index("Cost Media") + 1).value
        assert pct is None
        assert formula.startswith("=IF(") and '="",""' in formula


class TestDictionary:

    def test_unconfigured_slots_not_added(self, tmp_path):
        assert _roundtrip(tmp_path)["dictionary"] == {
            "lineitem_custom_dimensions": {
                "dim_custom1": {"status": "enabled", "caption": "Market"}
            },
        }


class TestLegacyWorkbook:
    """Workbooks exported by v3.0.12 and earlier must still import."""

    def test_old_labels_and_comma_lists(self, tmp_path):
        path = _export(tmp_path)
        wb = openpyxl.load_workbook(path)
        for sheet_name in ("Metadata", "Campaign"):
            for row in wb[sheet_name].iter_rows():
                if row[0].value == "Custom Properties (JSON):":
                    row[0].value = "Custom Properties:"
        locations = wb["Target Locations"]
        for cell in locations[1]:
            if cell.value and cell.value.endswith(" (JSON)"):
                cell.value = cell.value.replace(" (JSON)", "")
        locations.cell(row=2, column=4, value="New York, Chicago")
        wb.save(path)

        plan = import_from_excel(str(path))
        assert plan["meta"]["custom_properties"]["mt_plan_id"] == "P-1"
        assert plan["campaign"]["custom_properties"] == {"mt_client_id": "C-9"}
        assert plan["campaign"]["target_locations"][0]["location_list"] == ["New York", "Chicago"]


class TestFormulaCacheWarning:  # B7

    @pytest.mark.filterwarnings("default::mediaplanpy.excel.importer.ExcelFormulaCacheWarning")
    def test_warns_when_never_recalculated(self, tmp_path):
        with pytest.warns(ExcelFormulaCacheWarning, match="no cached value"):
            import_from_excel(str(_export(tmp_path)))

    @needs_soffice
    @pytest.mark.filterwarnings("error::mediaplanpy.excel.importer.ExcelFormulaCacheWarning")
    def test_silent_after_recalculation(self, tmp_path):
        # Recalculated workbooks contain blank formula results (B14), which
        # read back as None - they must not trigger the warning.
        import_from_excel(str(_recalculate(_export(tmp_path), tmp_path)))


class TestHelpers:

    @pytest.mark.parametrize(
        "value, expected",
        [
            ('["A", "B, C"]', ["A", "B, C"]),
            ("A, B", ["A", "B"]),
            ("", []),
            (None, []),
            ("[not json", ["[not json"]),
        ],
    )
    def test_parse_list_field(self, value, expected):
        assert _parse_list_field(value) == expected

    def test_merge_keeps_unrebuilt_and_extra_keys(self):
        existing = {
            "metric_impressions": {
                "formula_type": "cost_per_unit",
                "base_metric": "cost_total",
                "coefficient": 0.005,
                "comments": "note",
            },
            "metric_max_daily_spend": {"formula_type": "constant", "coefficient": 100},
        }
        rebuilt = {
            "metric_impressions": {
                "formula_type": "cost_per_unit",
                "base_metric": "cost_total",
                "coefficient": 0.006,
            },
            "metric_reach": {"formula_type": "constant", "base_metric": None, "coefficient": 5},
        }
        assert _merge_metric_formulas(existing, rebuilt) == {
            "metric_impressions": {
                "formula_type": "cost_per_unit",
                "base_metric": "cost_total",
                "coefficient": 0.006,
                "comments": "note",
            },
            "metric_max_daily_spend": {"formula_type": "constant", "coefficient": 100},
            "metric_reach": {"formula_type": "constant", "coefficient": 5},
        }
